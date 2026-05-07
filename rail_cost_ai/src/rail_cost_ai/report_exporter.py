from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document

from .models import ComparisonResult, CostItem, ProjectInfo


def export_excel_report(
    project: ProjectInfo,
    cost_items: list[CostItem],
    category_share_rows: list[dict],
    comparisons: list[ComparisonResult],
    output_file: str | Path,
) -> None:
    with pd.ExcelWriter(output_file, engine="xlsxwriter") as writer:
        pd.DataFrame([project.__dict__]).to_excel(writer, sheet_name="项目概览", index=False)
        pd.DataFrame([c.__dict__ for c in cost_items]).to_excel(writer, sheet_name="分项费用", index=False)
        pd.DataFrame(category_share_rows).to_excel(writer, sheet_name="占比分析", index=False)
        pd.DataFrame([c.__dict__ for c in comparisons]).to_excel(writer, sheet_name="历史对比", index=False)


def export_word_report(
    project: ProjectInfo,
    unit_cost: float,
    category_share_rows: list[dict],
    comparisons: list[ComparisonResult],
    output_file: str | Path,
) -> None:
    doc = Document()
    doc.add_heading("轨道交通造价分析报告", level=1)

    doc.add_heading("1. 项目基本信息", level=2)
    doc.add_paragraph(f"项目名称：{project.project_name}")
    doc.add_paragraph(f"线路长度：{project.line_length_km:.2f} 公里")
    doc.add_paragraph(f"车站数量：{project.station_count} 座")
    doc.add_paragraph(f"总投资：{project.total_investment_million:.2f} 百万元")
    doc.add_paragraph(f"单公里造价：{unit_cost:.2f} 百万元/公里")

    doc.add_heading("2. 分项费用占比", level=2)
    for row in category_share_rows:
        doc.add_paragraph(f"{row['category']}：{row['amount']:.2f} 百万元，占比 {row['share_pct']:.2f}%")

    doc.add_heading("3. 历史项目偏差分析", level=2)
    if not comparisons:
        doc.add_paragraph("暂无历史项目数据。")
    else:
        for item in comparisons[:5]:
            doc.add_paragraph(
                f"对比项目 {item.project_name}：历史单公里造价 {item.unit_cost_per_km:.2f}，偏差率 {item.deviation_rate_pct:.2f}%"
            )

    doc.save(output_file)
